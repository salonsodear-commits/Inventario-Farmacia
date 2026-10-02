# -*- coding: utf-8 -*-
"""Genera el archivo final: base SAP intacta + inventario fisico Avain cruzado.

Reglas de armado
----------------
* Las hojas SAP conservan columnas A:T, filas y formulas tal como venian.
* Se completa Q 'Cantidad fisica(recuento)' con el inventario fisico de Avain.
  Cuando un material tiene varias filas (lote/almacen), la cantidad fisica se
  imputa en la PRIMERA fila del material (logica BUSCARV) y las demas quedan en
  blanco: asi el total de la columna sigue siendo el inventario fisico real.
* Los items de Avain sin coincidencia en SAP se agregan como filas nuevas al
  final, sin codigo de material, para poder incorporar el inventario fisico.
* La fila de totales original NO se toca. Se agrega una fila TOTAL GENERAL al
  final con stock SAP, inventario fisico y diferencia.
* La informacion complementaria va en columnas nuevas a partir de U.
"""
import collections, json, copy, os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import load, match, paths

SRC = paths.dato(paths.STOCK_SAP)
OUT = os.path.join(paths.RAIZ, "Stock SAP vs Inventario Fisico Avain - cruce.xlsx")

# (primera fila de datos, ultima fila de datos, fila de totales original)
GEO = {"1120 SJ": (2, 679, 680), "1060 NQN": (2, 2023, 2024), "1130 Salta": (2, 485, 486)}
SHEET_BASE = {"1120 SJ": "San Juan", "1060 NQN": "Neuquén", "1130 Salta": "Salta"}
CENTRO = {"San Juan": "1120", "Neuquén": "1060", "Salta": "1130"}

C_MAT, C_TXT, C_CENTRO, C_ALM, C_UM, C_LIBRE = 1, 2, 6, 7, 8, 9
C_QFIS, C_DIFQ, C_DIFE, C_COM = 17, 18, 19, 20

EXTRA = [
    ("Nivel de coincidencia SAP–Avain", 26),
    ("Código Avain", 13),
    ("Nombre en inventario físico Avain", 42),
    ("Cantidad física Avain (ítem)", 14),
    ("Método de cruce", 34),
    ("¿Posible duplicidad SAP?", 13),
    ("Códigos SAP posibles equivalentes", 24),
    ("¿Informado en Master de Seba?", 14),
    ("ID común (Master)", 12),
    ("¿Aparece en Stock SAP?", 12),
    ("¿Aparece en inventario físico Avain?", 14),
    ("Stock SAP del material (todas las filas)", 15),
    ("Inventario físico Avain del material", 15),
    ("Diferencia a nivel material (SAP − físico)", 16),
    ("Observaciones del cruce", 90),
]
U0 = 21  # primera columna nueva

AZUL = "FF1F4E78"
GRIS = "FFF2F2F2"
FILL_EXTRA = PatternFill("solid", fgColor="FFDDEBF7")
FILL_NUEVA = PatternFill("solid", fgColor="FFFFF2CC")
FILL_TOTAL = PatternFill("solid", fgColor="FFE2EFDA")
NIVEL_FILL = {
    "Alta": PatternFill("solid", fgColor="FFC6EFCE"),
    "Media": PatternFill("solid", fgColor="FFFFEB9C"),
    "Baja": PatternFill("solid", fgColor="FFFCD5B4"),
    "Sin coincidencia": PatternFill("solid", fgColor="FFFFC7CE"),
}
THIN = Side(style="thin", color="FFBFBFBF")


def si_no(v):
    return "Sí" if v else "No"


def main():
    cruce = json.load(open(paths.salida("cruce.json"), encoding="utf-8"))
    sap = load.load_sap()
    wb = openpyxl.load_workbook(SRC, data_only=False)

    resumen = {}
    sin_coinc_global, dupes_global, fuera_master_global = [], [], []

    for sheet, (r0, r1, rtot) in GEO.items():
        base = SHEET_BASE[sheet]
        ws = wb[sheet]
        rows = sap[sheet]
        res = cruce[base]

        # --- agregacion por material: varios items Avain pueden ir al mismo material
        por_mat = collections.defaultdict(list)
        for rec in res:
            if rec["material"]:
                por_mat[rec["material"]].append(rec)

        filas_mat = collections.defaultdict(list)   # material -> filas de la hoja
        for r in rows:
            filas_mat[r["material"]].append(r["row"])
        stock_mat = collections.defaultdict(float)
        for r in rows:
            stock_mat[r["material"]] += (r["libre"] or 0)

        # --- encabezados de las columnas nuevas
        hdr_ref = ws.cell(1, 1)
        for i, (titulo, ancho) in enumerate(EXTRA):
            c = ws.cell(1, U0 + i, titulo)
            c.font = Font(name="Arial", sz=10, b=True, color="FFFFFFFF")
            c.fill = PatternFill("solid", fgColor=AZUL)
            c.alignment = Alignment(wrap_text=True, vertical="center",
                                    horizontal="center")
            c.border = Border(bottom=THIN)
            ws.column_dimensions[get_column_letter(U0 + i)].width = ancho

        # --- completar filas SAP existentes
        usados = set()
        for r in rows:
            fila = r["row"]
            mat = r["material"]
            recs = por_mat.get(mat, [])
            es_primera = filas_mat[mat][0] == fila
            fis_mat = sum(x["cantidad"] for x in recs)

            if recs and es_primera:
                ws.cell(fila, C_QFIS, fis_mat)
                usados.add(mat)

            # nivel resultante del material (el mas debil de los items imputados)
            nivel = ""
            if recs:
                nivel = recs[0]["nivel"]
                for x in recs[1:]:
                    nivel = match.peor(nivel, x["nivel"])

            obs = []
            dup, dupcods = False, []
            for x in recs:
                obs.extend(x["observaciones"])
                if x["sap_multiple"]:
                    dup = True
                    dupcods = x["sap_multiple"]
                if x["obs_origen"]:
                    obs.append(x["obs_origen"])
            if len(recs) > 1:
                obs.insert(0, "Varias denominaciones de Avain corresponden a este "
                              "material SAP (" +
                              "; ".join(f"{x['codigo_avain']} {x['nombre_avain']}"
                                        for x in recs) +
                              "). Se sumaron las cantidades físicas.")
            if recs and not es_primera:
                obs.insert(0, "La cantidad física del material se imputó en la fila "
                              f"{filas_mat[mat][0]} (primera fila del material).")
            if not recs:
                obs.append("Sin inventario físico de Avain asociado a este material.")

            vals = [
                nivel or "Sin inventario físico",
                "; ".join(x["codigo_avain"] for x in recs),
                "; ".join(x["nombre_avain"] for x in recs),
                (fis_mat if (recs and es_primera) else None),
                "; ".join(sorted({x["metodo"] for x in recs})),
                si_no(dup) if recs else "",
                ", ".join(dupcods),
                si_no(any(x["en_master"] for x in recs)) if recs
                else si_no(mat in _materiales_master()),
                "; ".join(sorted({x["id_master"] for x in recs if x["id_master"]})),
                "Sí",
                si_no(bool(recs)),
                stock_mat[mat],
                fis_mat if recs else 0,
                stock_mat[mat] - (fis_mat if recs else 0),
                " ".join(obs),
            ]
            for i, v in enumerate(vals):
                c = ws.cell(fila, U0 + i, v)
                c.font = Font(name="Arial", sz=10)
                c.fill = FILL_EXTRA
                if i in (3, 11, 12, 13):
                    c.number_format = "#,##0"
                if i == 14:
                    c.alignment = Alignment(wrap_text=False, vertical="top")
            ws.cell(fila, U0).fill = NIVEL_FILL.get(nivel, FILL_EXTRA)

        # --- items de Avain sin coincidencia -> filas nuevas al final
        sin = [x for x in res if not x["material"]]
        fila = rtot + 2
        ws.cell(fila, C_TXT,
                "— INVENTARIO FÍSICO AVAIN SIN COINCIDENCIA EN SAP "
                f"({len(sin)} ítems) — se agregan sin código de material para "
                "poder incorporar el inventario físico —")
        ws.cell(fila, C_TXT).font = Font(name="Arial", sz=10, b=True, color=AZUL)
        for cc in range(1, U0 + len(EXTRA)):
            ws.cell(fila, cc).fill = FILL_NUEVA
        fila += 1
        primera_nueva = fila
        for x in sorted(sin, key=lambda z: z["nombre_avain"].upper()):
            ws.cell(fila, C_MAT, None)                      # sin codigo SAP
            ws.cell(fila, C_TXT, x["nombre_avain"])
            ws.cell(fila, C_CENTRO, CENTRO[base])
            ws.cell(fila, C_UM, "C/U")
            ws.cell(fila, C_LIBRE, 0)
            ws.cell(fila, C_QFIS, x["cantidad"])
            ws.cell(fila, C_DIFQ, f"=I{fila}-Q{fila}")
            ws.cell(fila, C_COM,
                    "Ítem del inventario físico de Avain sin coincidencia en el "
                    "stock SAP de este centro. Se mantiene la línea sin código de "
                    "material; la diferencia indica una cantidad a incorporar al stock.")
            obs = list(x["observaciones"])
            if x["obs_origen"]:
                obs.append(x["obs_origen"])
            vals = ["Sin coincidencia", x["codigo_avain"], x["nombre_avain"],
                    x["cantidad"], x["metodo"], "No", "",
                    si_no(x["en_master"]), x["id_master"], "No", "Sí",
                    0, x["cantidad"], -x["cantidad"], " ".join(obs)]
            for i, v in enumerate(vals):
                c = ws.cell(fila, U0 + i, v)
                c.font = Font(name="Arial", sz=10)
                c.fill = FILL_EXTRA
                if i in (3, 11, 12, 13):
                    c.number_format = "#,##0"
            ws.cell(fila, U0).fill = NIVEL_FILL["Sin coincidencia"]
            for cc in (C_MAT, C_TXT, C_CENTRO, C_UM, C_LIBRE, C_QFIS, C_DIFQ, C_COM):
                ws.cell(fila, cc).font = Font(name="Arial", sz=10)
            ws.cell(fila, C_LIBRE).number_format = "#,##0"
            ws.cell(fila, C_QFIS).number_format = "#,##0"
            ws.cell(fila, C_DIFQ).number_format = "#,##0"
            sin_coinc_global.append({**x, "fila": fila, "hoja": sheet})
            fila += 1
        ultima = fila - 1

        # --- fila TOTAL GENERAL (la fila de totales original queda intacta)
        tg = ultima + 1
        ws.cell(tg, C_TXT, "TOTAL GENERAL — Stock SAP vs Inventario físico Avain")
        ws.cell(tg, C_LIBRE, f"=SUM(I{r0}:I{r1})+SUM(I{primera_nueva}:I{ultima})")
        ws.cell(tg, C_QFIS, f"=SUM(Q{r0}:Q{r1})+SUM(Q{primera_nueva}:Q{ultima})")
        ws.cell(tg, C_DIFQ, f"=I{tg}-Q{tg}")
        for cc in range(1, U0 + len(EXTRA)):
            c = ws.cell(tg, cc)
            c.fill = FILL_TOTAL
            c.font = Font(name="Arial", sz=10, b=True)
            c.border = Border(top=Side(style="medium", color=AZUL))
        for cc in (C_LIBRE, C_QFIS, C_DIFQ):
            ws.cell(tg, cc).number_format = "#,##0"
        ws.cell(tg, U0, "Diferencia = Stock SAP − Inventario físico Avain")

        ws.auto_filter.ref = f"A1:{get_column_letter(U0 + len(EXTRA) - 1)}{r1}"
        ws.freeze_panes = "C2"
        ws.row_dimensions[1].height = 58

        # --- resumen de la hoja
        c = collections.Counter()
        for r in rows:
            recs = por_mat.get(r["material"], [])
            if filas_mat[r["material"]][0] != r["row"]:
                continue
            if recs:
                niv = recs[0]["nivel"]
                for x in recs[1:]:
                    niv = match.peor(niv, x["nivel"])
                c[niv] += 1
            else:
                c["Sin inventario físico"] += 1
        con_fisico = {m for m, v in por_mat.items() if v}
        fis_por_mat = {m: sum(x["cantidad"] for x in v) for m, v in por_mat.items()}
        sap_con_fisico = sum(stock_mat[m] for m in con_fisico)
        a_bajar = sum(max(stock_mat[m] - fis_por_mat[m], 0) for m in con_fisico)
        a_sumar = sum(max(fis_por_mat[m] - stock_mat[m], 0) for m in con_fisico)
        sap_sin_fisico = sum(v for m, v in stock_mat.items() if m not in con_fisico)

        resumen[base] = {
            "hoja": sheet,
            "sap_con_fisico": sap_con_fisico,
            "sap_sin_fisico": sap_sin_fisico,
            "a_bajar": a_bajar,
            "a_sumar": a_sumar,
            "filas_sap": len(rows),
            "materiales_sap": len(filas_mat),
            "stock_sap": sum(stock_mat.values()),
            "items_avain": len(res),
            "fisico_total": sum(x["cantidad"] for x in res),
            "fisico_cruzado": sum(x["cantidad"] for x in res if x["material"]),
            "fisico_sin_coincidencia": sum(x["cantidad"] for x in res if not x["material"]),
            "items_sin_coincidencia": len(sin),
            "materiales_con_fisico": len(usados),
            "niveles_items": collections.Counter(x["nivel"] for x in res),
            "niveles_materiales": dict(c),
            "fila_total_original": rtot,
            "fila_total_general": tg,
            "primera_fila_nueva": primera_nueva,
            "ultima_fila_nueva": ultima,
        }
        for x in res:
            if x["sap_multiple"]:
                dupes_global.append({**x, "hoja": sheet})
            if x["material"] and not x["en_master"]:
                fuera_master_global.append({**x, "hoja": sheet})

    _hoja_resumen(wb, resumen)
    _hoja_lista(wb, "Sin coincidencia", sin_coinc_global,
                "Ítems del inventario físico de Avain sin coincidencia en el stock SAP. "
                "Se incorporaron como filas nuevas, sin código de material.")
    _hoja_lista(wb, "Posibles duplicidades", dupes_global,
                "El cruce SAP–Avain identificó más de un código SAP como posible "
                "equivalente del mismo ítem de Avain. Se imputó al primero "
                "(lógica BUSCARV). Requiere validación posterior.")
    _hoja_lista(wb, "Fuera del Master", fuera_master_global,
                "Ítems con stock físico y código SAP que NO estaban informados en el "
                "Master de Seba (la homologación cubre 284 ítems del master).")
    _hoja_metodologia(wb, resumen)

    wb.save(OUT)
    print("Generado:", OUT)
    for b, d in resumen.items():
        print(f"\n== {b} ({d['hoja']})")
        print(f"   stock SAP                {d['stock_sap']:>12,.0f}")
        print(f"   inventario físico Avain  {d['fisico_total']:>12,.0f}"
              f"  (cruzado {d['fisico_cruzado']:,.0f} / "
              f"sin coincidencia {d['fisico_sin_coincidencia']:,.0f})")
        print(f"   diferencia               "
              f"{d['stock_sap'] - d['fisico_total']:>12,.0f}")
        print(f"   niveles (ítems Avain):   {dict(d['niveles_items'])}")
        print(f"   filas nuevas             "
              f"{d['primera_fila_nueva']}–{d['ultima_fila_nueva']}"
              f" | TOTAL GENERAL fila {d['fila_total_general']}")
    json.dump({k: {kk: (dict(vv) if isinstance(vv, collections.Counter) else vv)
                   for kk, vv in v.items()} for k, v in resumen.items()},
              open(paths.salida("resumen.json"), "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)


_MM = None


def _materiales_master():
    global _MM
    if _MM is None:
        _MM = set()
        for h in load.load_homologacion():
            for c in h["sap_cod"]:
                _MM.add(c.strip())
    return _MM


def _titulo(ws, fila, texto, ancho=9):
    c = ws.cell(fila, 1, texto)
    c.font = Font(name="Arial", sz=12, b=True, color=AZUL)
    return fila + 1


def _encabezados(ws, fila, cols):
    for i, (t, w) in enumerate(cols, start=1):
        c = ws.cell(fila, i, t)
        c.font = Font(name="Arial", sz=10, b=True, color="FFFFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[fila].height = 32
    return fila + 1


def _hoja_resumen(wb, resumen):
    ws = wb.create_sheet("Resumen del cruce", 0)
    f = _titulo(ws, 2, "Cruce Stock SAP vs Inventario físico Avain")
    ws.cell(f, 1, "La base del archivo es el STOCK DE SAP y conserva su formato. "
                  "El inventario físico de Avain se cruzó sobre esa base. "
                  "La homologación se usó como vínculo de nomenclatura, no como base "
                  "de información.")
    ws.cell(f, 1).font = Font(name="Arial", sz=10, i=True)
    f += 2
    cols = [("Base / Centro", 16), ("Hoja SAP", 12), ("Stock SAP (Libre utilización)", 15),
            ("Inventario físico Avain", 15), ("Diferencia (SAP − físico)", 15),
            ("Stock SAP de materiales SÍ recontados", 16),
            ("Diferencia sobre lo recontado", 16),
            ("A dar de baja (SAP > físico)", 15),
            ("A incorporar (físico > SAP)", 15),
            ("Stock SAP sin recuento físico", 16),
            ("Físico cruzado a un código SAP", 15),
            ("Físico sin coincidencia", 14), ("Ítems Avain", 11),
            ("Alta", 8), ("Media", 8), ("Baja", 8), ("Sin coincidencia", 12),
            ("Materiales SAP", 12), ("Materiales con físico", 13)]
    f = _encabezados(ws, f, cols)
    tot = collections.Counter()
    for base, d in resumen.items():
        n = d["niveles_items"]
        vals = [base, d["hoja"], d["stock_sap"], d["fisico_total"],
                d["stock_sap"] - d["fisico_total"],
                d["sap_con_fisico"], d["sap_con_fisico"] - d["fisico_cruzado"],
                d["a_bajar"], d["a_sumar"], d["sap_sin_fisico"],
                d["fisico_cruzado"],
                d["fisico_sin_coincidencia"], d["items_avain"],
                n.get("Alta", 0), n.get("Media", 0), n.get("Baja", 0),
                n.get("Sin coincidencia", 0), d["materiales_sap"],
                d["materiales_con_fisico"]]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(f, i, v)
            c.font = Font(name="Arial", sz=10)
            if isinstance(v, (int, float)):
                c.number_format = "#,##0"
        for k in ("stock_sap", "fisico_total", "fisico_cruzado",
                  "fisico_sin_coincidencia", "items_avain", "sap_con_fisico",
                  "sap_sin_fisico", "a_bajar", "a_sumar"):
            tot[k] += d[k]
        for k in ("Alta", "Media", "Baja", "Sin coincidencia"):
            tot[k] += n.get(k, 0)
        f += 1
    vals = ["TOTAL", "", tot["stock_sap"], tot["fisico_total"],
            tot["stock_sap"] - tot["fisico_total"],
            tot["sap_con_fisico"], tot["sap_con_fisico"] - tot["fisico_cruzado"],
            tot["a_bajar"], tot["a_sumar"], tot["sap_sin_fisico"],
            tot["fisico_cruzado"],
            tot["fisico_sin_coincidencia"], tot["items_avain"],
            tot["Alta"], tot["Media"], tot["Baja"], tot["Sin coincidencia"], "", ""]
    for i, v in enumerate(vals, start=1):
        c = ws.cell(f, i, v)
        c.font = Font(name="Arial", sz=10, b=True)
        c.fill = FILL_TOTAL
        if isinstance(v, (int, float)):
            c.number_format = "#,##0"
    f += 3
    f = _titulo(ws, f, "Cómo leer la diferencia")
    for txt in [
        "Diferencia = Stock SAP (Libre utilización) − Inventario físico Avain, "
        "respetando la fórmula que ya traía el archivo SAP en la columna R.",
        "Diferencia POSITIVA: SAP informa más unidades que el recuento físico "
        "→ cantidad a dar de baja del stock.",
        "Diferencia NEGATIVA: el recuento físico supera lo informado por SAP "
        "→ cantidad a incorporar al stock. Es el caso esperado de los uniformes "
        "de San Juan.",
        "Los ítems sin coincidencia se agregaron al final de cada hoja sin código "
        "de material: su diferencia es negativa por definición (SAP = 0).",
    ]:
        c = ws.cell(f, 1, "• " + txt)
        c.font = Font(name="Arial", sz=10)
        f += 1
    ws.column_dimensions["A"].width = 18
    ws.sheet_view.showGridLines = False


LISTA_COLS = [("Hoja SAP", 12), ("Base", 12), ("Código Avain", 13),
              ("Nombre en inventario físico Avain", 44),
              ("Cantidad física", 13), ("Nivel", 18), ("Código SAP imputado", 15),
              ("Nomenclatura SAP", 40), ("Códigos SAP posibles equivalentes", 24),
              ("Método de cruce", 32), ("¿En Master de Seba?", 12),
              ("ID común (Master)", 12), ("Observaciones", 100)]


def _hoja_lista(wb, nombre, items, nota):
    ws = wb.create_sheet(nombre)
    f = _titulo(ws, 2, nombre)
    ws.cell(f, 1, nota)
    ws.cell(f, 1).font = Font(name="Arial", sz=10, i=True)
    f += 2
    f = _encabezados(ws, f, LISTA_COLS)
    for x in sorted(items, key=lambda z: (z["hoja"], -z["cantidad"])):
        vals = [x["hoja"], x["base"], x["codigo_avain"], x["nombre_avain"],
                x["cantidad"], x["nivel"], x["material"] or "",
                x["texto_sap"] or "", ", ".join(x["sap_multiple"]),
                x["metodo"], si_no(x["en_master"]), x["id_master"],
                " ".join(x["observaciones"])]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(f, i, v)
            c.font = Font(name="Arial", sz=10)
            if i == 5:
                c.number_format = "#,##0"
        ws.cell(f, 6).fill = NIVEL_FILL.get(x["nivel"], FILL_EXTRA)
        f += 1
    ws.auto_filter.ref = f"A5:{get_column_letter(len(LISTA_COLS))}{max(f-1,6)}"
    ws.freeze_panes = "A6"
    ws.sheet_view.showGridLines = False


def _hoja_metodologia(wb, resumen):
    ws = wb.create_sheet("Metodología")
    f = _titulo(ws, 2, "Metodología del cruce y criterios de certeza")
    bloques = [
        ("Base del proceso", [
            "La base es el archivo de STOCK DE SAP ('Stock Centro SJ NQN y Salta "
            "02.10.xlsx'). Se conservan sus hojas, columnas A:T, filas y fórmulas, "
            "porque el archivo lo usará Contabilidad.",
            "La homologación (Homologacion_Farmacia_AVAIN_con_SAP.xlsx) se usó "
            "únicamente como vínculo entre la nomenclatura SAP y los nombres de "
            "Avain. Cubre 284 ítems del Master, es decir información parcial: no es "
            "la base del inventario.",
            "El inventario físico se extrajo de los PDF enviados por Avain. "
            "San Juan: 202 renglones de lote (total declarado 202). "
            "Salta: 650 renglones (total declarado 650). "
            "Neuquén: informe IHSA, 243 ítems, 201 con stock (coincide con el "
            "encabezado del informe); se tomó la sección de Conciliación y se "
            "verificó contra las filas TOTAL MEDICAMENTO.",
        ]),
        ("Qué se completó en las hojas SAP", [
            "Columna Q 'Cantidad fisica(recuento)': inventario físico de Avain.",
            "Columna R 'Diferencia Q': fórmula original = Libre utilización − "
            "Cantidad física.",
            "Cuando un material SAP tiene varias filas (lote/almacén), la cantidad "
            "física se imputa en la PRIMERA fila del material y las demás quedan en "
            "blanco. Así el total de la columna es el inventario físico real y no se "
            "duplica. El detalle a nivel material está en las columnas "
            "complementarias.",
            "Los ítems de Avain sin coincidencia se agregaron como filas nuevas al "
            "final de cada hoja, sin código de material, con su cantidad física.",
            "La fila de totales original no se modificó. Se agregó una fila TOTAL "
            "GENERAL al final de cada hoja.",
        ]),
        ("Orden de precedencia del cruce", [
            "1. Homologación: código o nombre de Avain → ítem del Master → código "
            "SAP. La confianza resultante es la más débil de los dos tramos "
            "(Avain→Master y Master→SAP), porque el vínculo es una cadena.",
            "2. Nomenclatura: comparación del nombre de Avain contra el texto breve "
            "de material de SAP del mismo centro, con control de dosis, volumen y "
            "calibre/talle.",
            "3. Sin coincidencia: se conserva la línea y el código SAP queda vacío.",
        ]),
        ("Niveles de certeza", [
            "Alta: la nomenclatura coincide y la dosis/presentación es compatible, "
            "o la homologación respalda el vínculo con confianza alta en los dos "
            "tramos. Se cruza y se descuenta.",
            "Media: el ítem es el mismo pero falta un dato (una de las dos "
            "nomenclaturas no informa dosis o presentación), o la homologación "
            "aporta confianza media. Se cruza y se descuenta, y queda identificado.",
            "Baja: hay una incompatibilidad comprobada (dosis, volumen, calibre o "
            "talle distintos) o la evidencia es débil. Se cruza para no perder el "
            "recuento, pero debe validarse antes de tomarlo como definitivo.",
            "Sin coincidencia: no se identificó candidato respaldado en el stock SAP "
            "de ese centro. No prueba ausencia en todo SAP.",
        ]),
        ("Posibles duplicidades", [
            "Cuando dos o más códigos SAP pueden corresponder al mismo ítem de "
            "Avain, no se asume que sean iguales: se imputa el inventario al primer "
            "código (lógica BUSCARV/BUSCARX) y se deja la observación de que el "
            "cruce los identificó como posibles equivalentes del mismo ítem de "
            "Avain. Ver la hoja 'Posibles duplicidades'.",
            "Cuando varias denominaciones de Avain corresponden al mismo código SAP "
            "(por ejemplo el caso de la aspirineta, con tres nomenclaturas), se "
            "suman las cantidades físicas sobre ese material y se deja constancia "
            "en las observaciones.",
        ]),
        ("Calibración del criterio automático", [
            "El criterio de nomenclatura se contrastó contra las 267 filas del "
            "Master que ya tienen candidato SAP revisado por el equipo: reproduce "
            "el nivel exacto en el 56% de los casos y queda dentro de un nivel en "
            "el 79%. No produjo ninguna coincidencia Alta donde la revisión había "
            "dicho Baja o Sin coincidencia.",
            "Las diferencias restantes son casos de nombre comercial contra "
            "genérico (por ejemplo Meloxicam/FLEXIDOL, Dimenhidrinato/DRAMAMINE), "
            "que no se deducen del texto y sólo los resuelve la homologación.",
        ]),
        ("Hallazgos a validar", [
            "En el inventario de Neuquén hay 17 códigos usados para dos productos "
            "distintos (por ejemplo 529 = Tijera punta roma y Cloruro de potasio; "
            "532 = Mango de bisturí Nº 4 y Clonazepam 2 mg). El código no es clave "
            "única: el cruce se hizo por código + denominación.",
            "Dos ítems de Neuquén tienen diferencias internas en el propio informe "
            "de Avain (Adrenalina 1mg/mL: 15 unidades; Metronidazol 500mg: 20 "
            "unidades, entre stock actual y saldo de movimientos). Quedan marcados "
            "en las observaciones.",
            "La homologación tiene 89 ítems del Master con más de un código SAP "
            "candidato y 17 sin candidato. Resolverlos reduce el universo de "
            "niveles Media y Baja.",
        ]),
    ]
    for titulo, puntos in bloques:
        f += 1
        c = ws.cell(f, 1, titulo)
        c.font = Font(name="Arial", sz=11, b=True, color=AZUL)
        f += 1
        for p in puntos:
            c = ws.cell(f, 1, "• " + p)
            c.font = Font(name="Arial", sz=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[f].height = max(14, 13 * (1 + len(p) // 105))
            f += 1
    ws.column_dimensions["A"].width = 125
    ws.sheet_view.showGridLines = False


if __name__ == "__main__":
    main()
