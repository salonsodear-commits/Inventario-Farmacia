# -*- coding: utf-8 -*-
"""Estilos y utilidades de formato comunes del libro."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

AZUL = "FF1F4E78"
AZUL2 = "FF2E75B6"
THIN = Side(style="thin", color="FFBFBFBF")

F_BASE = Font(name="Arial", sz=10)
F_BOLD = Font(name="Arial", sz=10, b=True)
F_HDR = Font(name="Arial", sz=10, b=True, color="FFFFFFFF")
F_TIT = Font(name="Arial", sz=13, b=True, color=AZUL)
F_SUB = Font(name="Arial", sz=11, b=True, color=AZUL)
F_NOTA = Font(name="Arial", sz=9, i=True, color="FF595959")

FILL_HDR = PatternFill("solid", fgColor=AZUL)
FILL_HDR2 = PatternFill("solid", fgColor=AZUL2)
FILL_EXTRA = PatternFill("solid", fgColor="FFDDEBF7")
FILL_TOTAL = PatternFill("solid", fgColor="FFE2EFDA")
FILL_ALERTA = PatternFill("solid", fgColor="FFFFF2CC")
FILL_NUEVA = PatternFill("solid", fgColor="FFFCE4D6")

NIVEL_FILL = {
    "Alta": PatternFill("solid", fgColor="FFC6EFCE"),
    "Media": PatternFill("solid", fgColor="FFFFEB9C"),
    "Baja": PatternFill("solid", fgColor="FFFCD5B4"),
    "Sin coincidencia": PatternFill("solid", fgColor="FFFFC7CE"),
}
ACCION_FILL = {
    "Dar de baja en SAP": PatternFill("solid", fgColor="FFFCD5B4"),
    "Dar de alta en SAP": PatternFill("solid", fgColor="FFC6EFCE"),
    "Alta de material en SAP (sin código)": PatternFill("solid", fgColor="FFFFC7CE"),
    "Sin diferencia": PatternFill("solid", fgColor="FFF2F2F2"),
    "Sin recuento físico": PatternFill("solid", fgColor="FFEDEDED"),
    "Fuera del alcance del inventario físico": PatternFill("solid", fgColor="FFEDEDED"),
}
NUM = "#,##0"                      # UNIDADES
MON = '"$"\\ #,##0.00'             # IMPORTE (siempre con signo $)
L = get_column_letter


def titulo(ws, fila, texto, sub=None):
    c = ws.cell(fila, 1, texto)
    c.font = F_TIT
    fila += 1
    if sub:
        c = ws.cell(fila, 1, sub)
        c.font = F_NOTA
        c.alignment = Alignment(wrap_text=False, vertical="top")
        fila += 1
    return fila + 1


def encabezados(ws, fila, cols, fill=None):
    """cols: lista de (titulo, ancho). Devuelve la fila siguiente."""
    for i, (t, w) in enumerate(cols, start=1):
        c = ws.cell(fila, i, t)
        c.font = F_HDR
        c.fill = fill or FILL_HDR
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = Border(bottom=THIN)
        ws.column_dimensions[L(i)].width = w
    ws.row_dimensions[fila].height = 46
    return fila + 1


def subtitulo(ws, fila, texto):
    c = ws.cell(fila, 1, texto)
    c.font = F_SUB
    return fila + 1


def nota(ws, fila, texto, col=1):
    c = ws.cell(fila, col, texto)
    c.font = F_NOTA
    c.alignment = Alignment(wrap_text=False, vertical="top")
    return fila + 1
