# -*- coding: utf-8 -*-
"""Carga de las tres fuentes: SAP (base), inventarios Avain y homologacion."""
import json, re, collections
import openpyxl
import paths

SAP_FILE = paths.dato(paths.STOCK_SAP)
HOM_FILE = paths.HOMOLOGACION

# hoja SAP -> (primera fila de datos, ultima fila de datos, fila de totales, base)
SAP_SHEETS = {
    "1120 SJ":    (2, 679,  680, "San Juan"),
    "1060 NQN":   (2, 2023, 2024, "Neuquén"),
    "1130 Salta": (2, 485,  486, "Salta"),
}
# columnas SAP (1-based)
C_MAT, C_TXT, C_LOTE, C_GRP, C_TIPO, C_CENTRO, C_ALM, C_UM = 1, 2, 3, 4, 5, 6, 7, 8
C_LIBRE, C_TRANS, C_MON, C_VTRANS, C_VLIBRE, C_CTA = 9, 10, 11, 12, 13, 14
C_PCT, C_VU, C_QFIS, C_DIFQ, C_DIFE, C_COM = 15, 16, 17, 18, 19, 20
N_COLS = 20


def load_sap():
    """Filas SAP tal como vienen, por hoja."""
    wb = openpyxl.load_workbook(SAP_FILE, data_only=True)
    out = {}
    for sheet, (r0, r1, rtot, base) in SAP_SHEETS.items():
        ws = wb[sheet]
        rows = []
        for r in range(r0, r1 + 1):
            mat = ws.cell(r, C_MAT).value
            if mat is None or str(mat).strip() == "":
                continue
            rows.append({
                "sheet": sheet, "base": base, "row": r,
                "material": str(mat).strip(),
                "texto": (ws.cell(r, C_TXT).value or ""),
                "lote": ws.cell(r, C_LOTE).value,
                "grupo": ws.cell(r, C_GRP).value,
                "tipo": ws.cell(r, C_TIPO).value,
                "centro": ws.cell(r, C_CENTRO).value,
                "almacen": ws.cell(r, C_ALM).value,
                "um": ws.cell(r, C_UM).value,
                "libre": ws.cell(r, C_LIBRE).value or 0,
                "vu": ws.cell(r, C_VU).value,
                "vlibre": ws.cell(r, C_VLIBRE).value or 0,
            })
        out[sheet] = rows
    return out


def load_avain():
    """Inventarios fisicos Avain agregados por item (codigo + nombre)."""
    inv = {}
    for base, f in [("San Juan", "avain_San_Juan.json"), ("Salta", "avain_Salta.json")]:
        d = json.load(open(paths.salida(f), encoding="utf-8"))
        agg = {}
        for r in d["lotes"]:
            k = (r["codigo"], r["denominacion"])
            a = agg.setdefault(k, {"base": base, "codigo": r["codigo"],
                                   "nombre": r["denominacion"],
                                   "cantidad": 0.0, "lotes": 0, "obs": ""})
            a["cantidad"] += r["stock"]
            a["lotes"] += 1
        inv[base] = list(agg.values())
    d = json.load(open(paths.salida("avain_Neuquen.json"), encoding="utf-8"))
    inv["Neuquén"] = [{"base": "Neuquén", "codigo": r["codigo"], "nombre": r["medicamento"],
                       "cantidad": r["stock_actual"],
                       "lotes": None,
                       "obs": ("" if r["diagnostico"].startswith("Sin diferencias")
                               else "Informe Avain: " + r["diagnostico"])}
                      for r in d["conciliacion"]]
    return inv


def _split_multi(v):
    """Celda con varios valores separados por salto de linea."""
    if v is None:
        return []
    return [x.strip() for x in str(v).replace("\r", "").split("\n") if x.strip()]


# Catalogo comun: columnas por base
HOM_COLS = {
    "Neuquén":  {"cod": 4,  "nom": 5,  "conf": 6,  "clas": 7},
    "San Juan": {"cod": 8,  "nom": 9,  "conf": 10, "clas": 11},
    "Salta":    {"cod": 12, "nom": 13, "conf": 14, "clas": 15},
}
HOM_SAP = {"cod": 19, "nom": 20, "conf": 21, "clas": 22}
HOM_HDR_ROW, HOM_FIRST = 7, 8


def load_homologacion():
    """Catalogo comun -> filas del master con codigos por base y codigos SAP."""
    wb = openpyxl.load_workbook(HOM_FILE, data_only=True)
    ws = wb["Catálogo común"]
    rows = []
    for r in range(HOM_FIRST, ws.max_row + 1):
        mid = ws.cell(r, 1).value
        if not mid or not str(mid).strip():
            continue
        rec = {"id": str(mid).strip(), "tipo": ws.cell(r, 2).value,
               "nombre_master": ws.cell(r, 3).value, "fila": r,
               "sap_cod": _split_multi(ws.cell(r, HOM_SAP["cod"]).value),
               "sap_nom": _split_multi(ws.cell(r, HOM_SAP["nom"]).value),
               "sap_conf": ws.cell(r, HOM_SAP["conf"]).value,
               "sap_clas": ws.cell(r, HOM_SAP["clas"]).value,
               "bases": {}}
        for base, c in HOM_COLS.items():
            rec["bases"][base] = {
                "cod": _split_multi(ws.cell(r, c["cod"]).value),
                "nom": _split_multi(ws.cell(r, c["nom"]).value),
                "conf": ws.cell(r, c["conf"]).value,
                "clas": ws.cell(r, c["clas"]).value,
            }
        rows.append(rec)
    return rows


if __name__ == "__main__":
    sap = load_sap()
    for s, rows in sap.items():
        print(f"SAP {s}: {len(rows)} filas, {len({r['material'] for r in rows})} materiales, "
              f"libre={sum(r['libre'] for r in rows):,.0f}")
    inv = load_avain()
    for b, items in inv.items():
        print(f"Avain {b}: {len(items)} items, cantidad={sum(i['cantidad'] for i in items):,.0f}")
    hom = load_homologacion()
    print(f"Homologacion: {len(hom)} filas del master")
    nsap = sum(1 for h in hom if h["sap_cod"])
    multi = [h for h in hom if len(h["sap_cod"]) > 1]
    print(f"  con codigo SAP candidato: {nsap} | con MAS DE UN codigo SAP: {len(multi)}")
    print("  confianza SAP:", collections.Counter(h["sap_conf"] for h in hom))
    for h in multi[:5]:
        print("   *", h["id"], h["nombre_master"], "->", h["sap_cod"], h["sap_conf"])
