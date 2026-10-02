# -*- coding: utf-8 -*-
"""Verificacion del archivo generado: estructura, formulas y conciliacion."""
import json, re, collections, sys
import openpyxl
import load, paths, os

ARCH = os.path.join(paths.RAIZ, "Stock SAP vs Inventario Fisico Avain - cruce.xlsx")
GEO = {"1120 SJ": (2, 679, 680), "1060 NQN": (2, 2023, 2024), "1130 Salta": (2, 485, 486)}
BASE = {"1120 SJ": "San Juan", "1060 NQN": "Neuquén", "1130 Salta": "Salta"}
COL = {"I": 9, "Q": 17, "R": 18}

ok_all = True


def check(cond, msg):
    global ok_all
    print(("  OK   " if cond else "  FALLA") + "  " + msg)
    if not cond:
        ok_all = False


def suma_rango(ws, col, a, b):
    t = 0.0
    for r in range(a, b + 1):
        v = ws.cell(r, COL[col]).value
        if isinstance(v, (int, float)):
            t += v
    return t


def eval_sumas(ws, formula):
    """Evalua '=SUM(I2:I679)+SUM(I683:I734)'."""
    tot = 0.0
    for m in re.finditer(r"SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)", formula.replace("$", "")):
        c, a, _, b = m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
        tot += suma_rango(ws, c, a, b)
    return tot


def main():
    cruce = json.load(open(paths.salida("cruce.json"), encoding="utf-8"))
    sap = load.load_sap()
    wb = openpyxl.load_workbook(ARCH, data_only=False)

    print("HOJAS:", wb.sheetnames)
    check(all(s in wb.sheetnames for s in GEO), "las tres hojas SAP se conservan")
    check(wb.sheetnames.index("1120 SJ") < wb.sheetnames.index("1060 NQN")
          < wb.sheetnames.index("1130 Salta"), "orden original de las hojas SAP")

    for sheet, (r0, r1, rtot) in GEO.items():
        print(f"\n===== {sheet}")
        ws = wb[sheet]
        base = BASE[sheet]
        rows = sap[sheet]
        res = cruce[base]

        # --- 1. la base SAP quedo intacta
        difs = 0
        for r in rows:
            if str(ws.cell(r["row"], 1).value).strip() != r["material"]:
                difs += 1
            if (ws.cell(r["row"], 9).value or 0) != (r["libre"] or 0):
                difs += 1
        check(difs == 0, f"columnas Material y Libre utilización intactas en {len(rows)} filas")
        check(ws["R2"].value == "=I2-Q2", "fórmula original de Diferencia Q intacta (R2)")
        check(ws["S2"].value == "=P2*Q2", "fórmula original de Diferencia económica intacta (S2)")
        check(str(ws[f"O2"].value).endswith(f"$M${rtot}"),
              f"fórmula O2 sigue apuntando a $M${rtot} (fila de totales original sin mover)")
        check(ws.cell(rtot, 13).value == f"=SUM(M2:M{r1})",
              "fila de totales original conservada")

        # --- 2. una sola imputacion de cantidad fisica por material
        filas_mat = collections.defaultdict(list)
        for r in rows:
            filas_mat[r["material"]].append(r["row"])
        fis = collections.defaultdict(float)
        for x in res:
            if x["material"]:
                fis[x["material"]] += x["cantidad"]
        malos = []
        for mat, qty in fis.items():
            con_valor = [f for f in filas_mat[mat]
                         if isinstance(ws.cell(f, 17).value, (int, float))]
            if len(con_valor) != 1 or ws.cell(con_valor[0], 17).value != qty:
                malos.append(mat)
        check(not malos, f"cada material con recuento tiene la cantidad en UNA sola fila "
                         f"({len(fis)} materiales){'' if not malos else ' -> ' + str(malos[:5])}")

        # --- 3. filas nuevas
        nuevas = [r for r in range(rtot + 1, ws.max_row + 1)
                  if isinstance(ws.cell(r, 17).value, (int, float))
                  and ws.cell(r, 2).value and not str(ws.cell(r, 2).value).startswith("TOTAL")]
        sin = [x for x in res if not x["material"]]
        check(len(nuevas) == len(sin),
              f"filas nuevas agregadas = ítems sin coincidencia ({len(nuevas)} = {len(sin)})")
        check(all(ws.cell(r, 1).value in (None, "") for r in nuevas),
              "las filas nuevas no llevan código de material (columna A vacía)")
        check(all(ws.cell(r, 18).value == f"=I{r}-Q{r}" for r in nuevas),
              "las filas nuevas llevan la fórmula de diferencia")

        # --- 4. conciliacion de totales
        tg = ws.max_row
        while tg > 1 and not str(ws.cell(tg, 2).value or "").startswith("TOTAL GENERAL"):
            tg -= 1
        check(tg > rtot, f"fila TOTAL GENERAL presente (fila {tg})")
        i_tot = eval_sumas(ws, ws.cell(tg, 9).value)
        q_tot = eval_sumas(ws, ws.cell(tg, 17).value)
        sap_esp = sum(r["libre"] or 0 for r in rows)
        fis_esp = sum(x["cantidad"] for x in res)
        check(abs(i_tot - sap_esp) < 1e-6,
              f"TOTAL GENERAL stock SAP = {i_tot:,.0f} (esperado {sap_esp:,.0f})")
        check(abs(q_tot - fis_esp) < 1e-6,
              f"TOTAL GENERAL inventario físico = {q_tot:,.0f} (esperado {fis_esp:,.0f})")
        print(f"       diferencia TOTAL GENERAL = {i_tot - q_tot:,.0f}")

        # --- 5. columnas complementarias
        hdr = [ws.cell(1, c).value for c in range(21, 36)]
        check(hdr[0] == "Nivel de coincidencia SAP–Avain" and hdr[-1] == "Observaciones del cruce",
              "columnas complementarias presentes de U a AI")
        niveles = collections.Counter(ws.cell(r, 21).value for r in range(r0, r1 + 1))
        print("       niveles por fila SAP:", dict(niveles))

    # --- hojas auxiliares
    print("\n===== hojas auxiliares")
    for nombre, esperado in [("Sin coincidencia", sum(len([x for x in cruce[b] if not x["material"]]) for b in cruce)),
                             ("Posibles duplicidades", sum(len([x for x in cruce[b] if x["sap_multiple"]]) for b in cruce))]:
        ws = wb[nombre]
        n = sum(1 for r in range(6, ws.max_row + 1) if ws.cell(r, 3).value or ws.cell(r, 4).value)
        check(n == esperado, f"hoja '{nombre}': {n} filas (esperado {esperado})")

    print("\n" + ("TODAS LAS VERIFICACIONES OK" if ok_all else "HAY VERIFICACIONES EN FALLA"))
    return 0 if ok_all else 1


if __name__ == "__main__":
    sys.exit(main())
