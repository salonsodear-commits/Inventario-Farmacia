# -*- coding: utf-8 -*-
"""Control de cantidades contra las fuentes originales (PDF de Avain y SAP)."""
import json, collections
from openpyxl.styles import Border, Side
import estilo as E
from estilo import L
import paths

COLS = [("Fuente", 26), ("Base", 12), ("Concepto", 46),
        ("Total en la fuente original", 17), ("Total usado en el análisis", 17),
        ("Diferencia", 13), ("Estado", 14), ("Detalle", 95)]


def _avain():
    """Totales de los PDF, recalculados desde la extraccion."""
    out = {}
    for base, f in [("San Juan", "avain_San_Juan.json"),
                    ("Salta", "avain_Salta.json")]:
        d = json.load(open(paths.salida(f), encoding="utf-8"))
        out[base] = {
            "suma_stock_pdf": round(sum(r["stock"] for r in d["lotes"]), 2),
            "renglones": d["control"]["filas_lote"],
            "renglones_declarados": d["control"].get("total_declarado"),
            "items": len(d["items"]),
            "suma_items": round(sum(i["stock"] for i in d["items"]), 2),
        }
    d = json.load(open(paths.salida("avain_Neuquen.json"), encoding="utf-8"))
    con = d["conciliacion"]
    out["Neuquén"] = {
        "suma_stock_pdf": round(sum(r["stock_actual"] for r in con), 2),
        "suma_totales_seccion": round(sum(r["stock_actual"] for r in d["totales"]), 2),
        "suma_saldos_lote": round(sum(r["saldo_lote"] for r in d["detalle"]), 2),
        "items": len(con),
        "con_stock": sum(1 for r in con if r["stock_actual"] > 0),
        "declarado_con_stock": d["meta"].get("con_stock_declarado"),
        "suma_items": round(sum(r["stock_actual"] for r in con), 2),
    }
    return out


def escribir(wb, filas_conc, sap, geo):
    ws = wb.create_sheet("Control de cantidades")
    f = E.titulo(ws, 2, "Control de cantidades contra las fuentes originales",
                 "La columna «Total en la fuente original» es el dato de los "
                 "archivos originales. La columna «Total usado» es una FÓRMULA VIVA "
                 "sobre las hojas de detalle: si el libro pierde o cambia una fila, "
                 "el estado pasa a A REVISAR automáticamente. Todo debe decir "
                 "Coincide salvo la última línea, que es una diferencia del propio "
                 "informe de origen.")
    f = E.encabezados(ws, f, COLS)
    ini = f
    av = _avain()
    SH = {"San Juan": "1120 SJ", "Neuquén": "1060 NQN", "Salta": "1130 Salta"}

    def linea(fuente, base, concepto, orig, usado, detalle, formato=E.NUM):
        """usado puede ser un valor o una formula: si es formula, el control es
        vivo y detecta cualquier fila que se pierda despues de generado."""
        nonlocal f
        for i, v in enumerate([fuente, base, concepto, orig, usado], start=1):
            c = ws.cell(f, i, v)
            c.font = E.F_BASE
            if i in (4, 5):
                c.number_format = formato
        ws.cell(f, 6, f"=D{f}-E{f}").number_format = formato
        ws.cell(f, 6).font = E.F_BASE
        ws.cell(f, 7, f'=IF(ABS(F{f})<0.005,"Coincide","A REVISAR")')
        ws.cell(f, 7).font = E.F_BOLD
        ws.cell(f, 8, detalle).font = E.F_BASE
        f += 1

    # ---- cantidad de filas de SAP: detecta al instante una fila perdida
    for base, sh in SH.items():
        g = geo[sh]
        linea(f"Bajada SAP · hoja {sh}", base,
              "Cantidad de filas de material de SAP",
              len(sap[sh]),
              f"=COUNTA('{sh}'!$A${g['r0']}:$A${g['r1']})",
              "Si este control dice A REVISAR, la hoja perdió o ganó filas "
              "respecto de la bajada original y el resto de los números no es "
              "confiable.")

    # ---- stock SAP (formula viva sobre la hoja de detalle)
    for base, sh in SH.items():
        g = geo[sh]
        linea(f"Bajada SAP · hoja {sh}", base,
              "Suma de la columna Libre utilización",
              sum(r["libre"] or 0 for r in sap[sh]),
              f"=SUM('{sh}'!$I${g['r0']}:$I${g['r1']})",
              f"{len(sap[sh])} filas de SAP (material / lote / almacén) agrupadas "
              f"en {len({r['material'] for r in sap[sh]})} materiales. La hoja de "
              f"detalle conserva las filas originales sin modificar.")
    for base, sh in SH.items():
        g = geo[sh]
        linea(f"Bajada SAP · hoja {sh}", base,
              "Suma de la columna Valor libre util.",
              round(sum(r["vlibre"] or 0 for r in sap[sh]), 2),
              f"=SUM('{sh}'!$M${g['r0']}:$M${g['r1']})",
              "Control de la valorización contra el importe informado por SAP.",
              formato=E.MON)

    # ---- inventarios fisicos (formula viva sobre la columna del recuento)
    for base, sh in SH.items():
        g = geo[sh]
        a = av[base]
        if base == "Neuquén":
            det = (f"{a['items']} ítems, {a['con_stock']} con stock, que coincide "
                   f"con el encabezado del informe («Con stock: "
                   f"{a['declarado_con_stock']} medicamentos»).")
        else:
            det = (f"{a['renglones']} renglones de lote agrupados en {a['items']} "
                   f"ítems. El rótulo «TOTAL: {a['renglones_declarados']:.0f}» del "
                   f"PDF es la cantidad de renglones, no la suma de stock.")
        linea(f"Inventario físico · {base}", base,
              "Suma del recuento físico informado en el PDF",
              a["suma_stock_pdf"],
              f"=SUM('{sh}'!$Q${g['r0']}:$Q${g['ultima']})", det)
    a = av["Neuquén"]
    g = geo["1060 NQN"]
    linea("Inventario físico · Neuquén", "Neuquén",
          "Suma de los saldos por lote del informe",
          a["suma_saldos_lote"],
          f"=SUM('1060 NQN'!$Q${g['r0']}:$Q${g['ultima']})",
          "Diferencia propia del informe, no del cruce: Adrenalina 1mg/mL (15) y "
          "Metronidazol 500mg (20) difieren entre el stock actual y el saldo de "
          "movimientos. Se conservó el stock actual, que es el total declarado.")
    fin = f - 1
    for cc in range(1, len(COLS) + 1):
        ws.cell(f, cc).border = Border(top=Side(style="medium", color=E.AZUL))
    ws.auto_filter.ref = f"A{ini-1}:{L(len(COLS))}{fin}"
    ws.freeze_panes = f"A{ini}"
    ws.sheet_view.showGridLines = False
    return av
